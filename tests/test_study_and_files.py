from __future__ import annotations

import io
import tempfile
from pathlib import Path

from fastapi.testclient import TestClient

from aegismind_core.app import create_app
from aegismind_core.routes import CoreState, extract_text_from_file_bytes


def test_extract_text_plain_text() -> None:
    content = "Hello AegisMind, this is plain text content."
    text, file_type, count = extract_text_from_file_bytes("sample.txt", content.encode("utf-8"))
    assert "Hello AegisMind" in text
    assert file_type == "txt"
    assert count == 1


def test_extract_text_docx() -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("Architecture Overview", 0)
    doc.add_paragraph("AegisMind sovereign retrieval system.")
    buf = io.BytesIO()
    doc.save(buf)

    text, file_type, count = extract_text_from_file_bytes("overview.docx", buf.getvalue())
    assert "Architecture Overview" in text
    assert "sovereign retrieval system" in text
    assert file_type == "document"
    assert count == 1


def test_extract_text_pptx() -> None:
    import pptx

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Key Slides on Sovereign Access"
    buf = io.BytesIO()
    prs.save(buf)

    text, file_type, count = extract_text_from_file_bytes("lecture.pptx", buf.getvalue())
    assert "Key Slides on Sovereign Access" in text
    assert file_type == "presentation"
    assert count == 1


def test_study_ask_and_notes_flow() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        notes_dir = Path(tmpdir) / "notes"
        notes_dir.mkdir()

        app = create_app(state=CoreState())
        client = TestClient(app)

        # 1. Parse File Endpoint
        file_bytes = (
            b"Course Lecture 1:\nIntroduction to Vector Search and Sovereign Access Control."
        )
        parse_res = client.post(
            "/api/v1/documents/parse-file",
            files={"file": ("lecture1.txt", file_bytes, "text/plain")},
        )
        assert parse_res.status_code == 200
        parsed = parse_res.json()
        assert parsed["title"] == "Lecture1"
        assert "Vector Search" in parsed["content"]

        # 2. Study Ask Endpoint: Q&A Mode
        qa_res = client.post(
            "/api/v1/study/ask",
            json={
                "title": parsed["title"],
                "content": parsed["content"],
                "query": "What are the main topics?",
                "mode": "qa",
            },
        )
        assert qa_res.status_code == 200
        assert "answer" in qa_res.json()

        # 3. Study Ask Endpoint: Quiz Mode
        quiz_res = client.post(
            "/api/v1/study/ask",
            json={
                "title": parsed["title"],
                "content": parsed["content"],
                "query": "Quiz me on this lecture",
                "mode": "quiz",
            },
        )
        assert quiz_res.status_code == 200
        quiz_answer = quiz_res.json()["answer"].lower()
        assert "quiz" in quiz_answer or "question" in quiz_answer

        # 4. Study Ask Endpoint: Summary Mode
        sum_res = client.post(
            "/api/v1/study/ask",
            json={
                "title": parsed["title"],
                "content": parsed["content"],
                "query": "Summarize key points",
                "mode": "summary",
            },
        )
        assert sum_res.status_code == 200
        assert len(sum_res.json()["answer"]) > 10

        # 5. Create Note in Vault
        note_res = client.post(
            "/api/v1/notes",
            json={
                "title": "Study Notes: Lecture 1",
                "content": "Key takeaways on Vector Search and Sovereign Access.",
                "tags": ["study", "lecture"],
                "notes_dir": str(notes_dir),
            },
        )
        assert note_res.status_code == 200
        assert note_res.json()["status"] == "created"

        # Verify note was created and listed
        list_res = client.get(f"/api/v1/notes?notes_dir={notes_dir}")
        assert list_res.status_code == 200
        notes = list_res.json()
        assert len(notes) == 1
        assert "Study Notes: Lecture 1" in notes[0]["title"]
